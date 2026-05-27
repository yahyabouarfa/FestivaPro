from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import text
from sqlalchemy.orm import Session


@dataclass(frozen=True)
class ReportFile:
    filename: str
    content_type: str
    data: bytes


def money(value) -> Decimal:
    return Decimal(value or 0).quantize(Decimal("0.01"))


def fetch_event(db: Session, event_id: int):
    return db.execute(text("SELECT * FROM events WHERE id = :event_id"), {"event_id": event_id}).mappings().first()


def fetch_bar(db: Session, bar_id: int):
    return db.execute(
        text(
            """
            SELECT b.*, e.name AS event_name, e.event_date
            FROM bars b
            JOIN events e ON e.id = b.event_id
            WHERE b.id = :bar_id
            """
        ),
        {"bar_id": bar_id},
    ).mappings().first()


def bar_products(db: Session, bar_id: int):
    return db.execute(
        text(
            """
            SELECT p.name AS product_name, pc.name AS category_name, bs.quantity_allocated, bs.quantity_sold,
                   bs.quantity_remaining, f.revenue, f.cost, f.profit
            FROM bar_stock bs
            JOIN products p ON p.id = bs.product_id
            JOIN product_categories pc ON pc.id = p.category_id
            JOIN bar_stock_financials f ON f.bar_stock_id = bs.id
            WHERE bs.bar_id = :bar_id
            ORDER BY pc.name, p.name
            """
        ),
        {"bar_id": bar_id},
    ).mappings().all()


def bar_bartenders(db: Session, bar_id: int):
    return db.execute(
        text(
            """
            SELECT u.full_name, COALESCE(s.units_sold, 0) AS units_sold,
                   COALESCE(s.contribution_pct, 0) AS contribution_pct,
                   COALESCE(es.salary_amount, 300) AS salary_paid
            FROM event_salaries es
            JOIN bars b ON b.id = es.bar_id
            JOIN users u ON u.id = es.user_id
            LEFT JOIN bartender_sales s ON s.event_id = es.event_id AND s.bar_id = es.bar_id AND s.user_id = es.user_id
            WHERE es.bar_id = :bar_id
            ORDER BY u.full_name
            """
        ),
        {"bar_id": bar_id},
    ).mappings().all()


def bars_summary(db: Session, event_id: int):
    return db.execute(
        text(
            """
            SELECT b.id AS bar_id, b.name AS bar_name,
                   COALESCE(SUM(f.quantity_allocated), 0) AS allocated,
                   COALESCE(SUM(f.quantity_sold), 0) AS sold,
                   COALESCE(SUM(f.quantity_remaining), 0) AS remaining,
                   COALESCE(SUM(f.revenue), 0) AS revenue,
                   COALESCE(SUM(f.cost), 0) AS cost,
                   COALESCE(payroll.staff_cost, 0) AS staff_cost,
                   COALESCE(SUM(f.profit), 0) - COALESCE(payroll.staff_cost, 0) AS net_profit,
                   CASE WHEN COALESCE(SUM(f.quantity_allocated), 0) = 0 THEN 0
                        ELSE COALESCE(SUM(f.quantity_remaining), 0) / SUM(f.quantity_allocated) * 100 END AS waste_pct
            FROM bars b
            LEFT JOIN bar_stock_financials f ON f.bar_id = b.id
            LEFT JOIN (
              SELECT bar_id, SUM(salary_amount) AS staff_cost
              FROM event_salaries
              GROUP BY bar_id
            ) payroll ON payroll.bar_id = b.id
            WHERE b.event_id = :event_id
            GROUP BY b.id, b.name, payroll.staff_cost
            ORDER BY net_profit DESC, revenue DESC, waste_pct ASC
            """
        ),
        {"event_id": event_id},
    ).mappings().all()


def top_products(db: Session, event_id: int):
    return db.execute(
        text(
            """
            SELECT p.name AS product_name, SUM(f.quantity_sold) AS units_sold, SUM(f.revenue) AS revenue
            FROM bar_stock_financials f
            JOIN products p ON p.id = f.product_id
            WHERE f.event_id = :event_id
            GROUP BY p.id, p.name
            ORDER BY revenue DESC
            LIMIT 10
            """
        ),
        {"event_id": event_id},
    ).mappings().all()


def payroll(db: Session, event_id: int):
    return db.execute(
        text(
            """
            SELECT u.full_name, b.name AS bar_name, SUM(es.salary_amount) AS salary_paid
            FROM event_salaries es
            JOIN users u ON u.id = es.user_id
            JOIN bars b ON b.id = es.bar_id
            WHERE es.event_id = :event_id
            GROUP BY u.id, u.full_name, b.name
            ORDER BY u.full_name
            """
        ),
        {"event_id": event_id},
    ).mappings().all()


def financial_totals(rows):
    revenue = sum(money(row["revenue"]) for row in rows)
    cost = sum(money(row["cost"]) for row in rows)
    staff = sum(money(row["staff_cost"]) for row in rows)
    return {"revenue": revenue, "cost": cost, "staff_cost": staff, "net_profit": revenue - cost - staff}


def report_signature_table(labels: list[str], styles) -> Table:
    table = Table(
        [
            [Paragraph(f"<b>{label}</b>", styles["Normal"]) for label in labels],
            ["Nom: ____________________" for _ in labels],
            ["Signature: _______________" for _ in labels],
            ["Date: ____________________" for _ in labels],
        ],
        colWidths=[240 if len(labels) == 2 else 160 for _ in labels],
    )
    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTSIZE", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    return table


def build_pdf(title: str, sections: list[tuple[str, list[str], list[list[object]]]], signature_mode: str = "general") -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("<b>FestivaPro</b> | Festival / Nightlife Operations", styles["Title"]),
        Paragraph(title, styles["Heading2"]),
        Spacer(1, 10),
    ]
    for heading, headers, rows in sections:
        story.append(Paragraph(heading, styles["Heading3"]))
        table = Table([headers, *rows], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#171124")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#f8f7ff")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d8d2e8")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f1ff")]),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.extend([table, Spacer(1, 12)])
    labels = ["Responsable du stock", "Directeur"]
    if signature_mode == "bar":
        labels = ["Responsable du bar", "Responsable du stock", "Directeur"]
    story.append(report_signature_table(labels, styles))
    doc.build(story)
    return buffer.getvalue()


def workbook_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def add_sheet(wb: Workbook, title: str, headers: list[str], rows) -> None:
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)


def per_bar_report(db: Session, bar_id: int, file_type: str) -> ReportFile:
    bar = fetch_bar(db, bar_id)
    products = bar_products(db, bar_id)
    bartenders = bar_bartenders(db, bar_id)
    staff_cost = sum(money(row["salary_paid"]) for row in bartenders)
    revenue = sum(money(row["revenue"]) for row in products)
    cost = sum(money(row["cost"]) for row in products)
    totals = [["Gross revenue", revenue], ["COGS", cost], ["Staff cost", staff_cost], ["Net profit", revenue - cost - staff_cost]]
    base = f"bar-{bar_id}-night-report"

    if file_type == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        add_sheet(wb, "Overview", ["Field", "Value"], [["Bar", bar["name"]], ["Event", bar["event_name"]], ["Date", str(bar["event_date"])], *totals])
        add_sheet(wb, "Products", ["Product", "Category", "Allocated", "Sold", "Remaining", "Revenue", "Cost", "Profit"], [[p["product_name"], p["category_name"], p["quantity_allocated"], p["quantity_sold"], p["quantity_remaining"], p["revenue"], p["cost"], p["profit"]] for p in products])
        add_sheet(wb, "Bartenders", ["Bartender", "Units sold", "Contribution %", "Salary paid"], [[b["full_name"], b["units_sold"], b["contribution_pct"], b["salary_paid"]] for b in bartenders])
        return ReportFile(f"{base}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", workbook_bytes(wb))

    pdf = build_pdf(
        f"{bar['name']} | {bar['event_name']} | {bar['event_date']}",
        [
            ("Bar Overview", ["Field", "Value"], totals),
            ("Product Breakdown", ["Product", "Category", "Allocated", "Sold", "Remaining", "Revenue", "Cost", "Profit"], [[p["product_name"], p["category_name"], p["quantity_allocated"], p["quantity_sold"], p["quantity_remaining"], p["revenue"], p["cost"], p["profit"]] for p in products]),
            ("Bartenders", ["Bartender", "Units", "Contribution %", "Salary"], [[b["full_name"], b["units_sold"], b["contribution_pct"], b["salary_paid"]] for b in bartenders]),
        ],
        signature_mode="bar",
    )
    return ReportFile(f"{base}.pdf", "application/pdf", pdf)


def all_bars_report(db: Session, event_id: int, file_type: str) -> ReportFile:
    event = fetch_event(db, event_id)
    rows = bars_summary(db, event_id)
    totals = financial_totals(rows)
    data = [[r["bar_name"], r["sold"], r["remaining"], r["revenue"], r["cost"], r["staff_cost"], r["net_profit"], r["waste_pct"]] for r in rows]
    base = f"event-{event_id}-all-bars-report"
    if file_type == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        add_sheet(wb, "All Bars", ["Bar", "Sold", "Remaining", "Revenue", "Cost", "Staff", "Net Profit", "Waste %"], data)
        add_sheet(wb, "Totals", ["Metric", "Value"], [["Revenue", totals["revenue"]], ["COGS", totals["cost"]], ["Staff cost", totals["staff_cost"]], ["Net profit", totals["net_profit"]]])
        return ReportFile(f"{base}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", workbook_bytes(wb))
    return ReportFile(
        f"{base}.pdf",
        "application/pdf",
        build_pdf(f"All Bars | {event['name']} | {event['event_date']}", [("Ranked Bar Summary", ["Bar", "Sold", "Remaining", "Revenue", "Cost", "Staff", "Net Profit", "Waste %"], data), ("Totals", ["Metric", "Value"], [["Revenue", totals["revenue"]], ["COGS", totals["cost"]], ["Staff cost", totals["staff_cost"]], ["Net profit", totals["net_profit"]]])]),
    )


def full_event_report(db: Session, event_id: int, file_type: str) -> ReportFile:
    event = fetch_event(db, event_id)
    bar_rows = bars_summary(db, event_id)
    product_rows = top_products(db, event_id)
    payroll_rows = payroll(db, event_id)
    totals = financial_totals(bar_rows)
    base = f"event-{event_id}-full-report"
    if file_type == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        add_sheet(wb, "P&L", ["Metric", "Value"], [["Revenue", totals["revenue"]], ["COGS", totals["cost"]], ["Payroll", totals["staff_cost"]], ["Net profit", totals["net_profit"]]])
        add_sheet(wb, "Bars", ["Bar", "Revenue", "Cost", "Staff", "Net Profit", "Waste %"], [[r["bar_name"], r["revenue"], r["cost"], r["staff_cost"], r["net_profit"], r["waste_pct"]] for r in bar_rows])
        add_sheet(wb, "Top Products", ["Product", "Units sold", "Revenue"], [[r["product_name"], r["units_sold"], r["revenue"]] for r in product_rows])
        add_sheet(wb, "Payroll", ["Employee", "Bar", "Salary"], [[r["full_name"], r["bar_name"], r["salary_paid"]] for r in payroll_rows])
        bars_ws = wb["Bars"]
        chart = BarChart()
        chart.title = "Revenue by bar"
        chart.add_data(Reference(bars_ws, min_col=2, min_row=1, max_row=bars_ws.max_row), titles_from_data=True)
        chart.set_categories(Reference(bars_ws, min_col=1, min_row=2, max_row=bars_ws.max_row))
        bars_ws.add_chart(chart, "H2")
        products_ws = wb["Top Products"]
        pie = PieChart()
        pie.title = "Product sales mix"
        pie.add_data(Reference(products_ws, min_col=2, min_row=1, max_row=products_ws.max_row), titles_from_data=True)
        pie.set_categories(Reference(products_ws, min_col=1, min_row=2, max_row=products_ws.max_row))
        products_ws.add_chart(pie, "E2")
        return ReportFile(f"{base}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", workbook_bytes(wb))
    return ReportFile(
        f"{base}.pdf",
        "application/pdf",
        build_pdf(
            f"Full Event Report | {event['name']} | {event['event_date']}",
            [
                ("Final P&L", ["Metric", "Value"], [["Revenue", totals["revenue"]], ["COGS", totals["cost"]], ["Payroll", totals["staff_cost"]], ["Net profit", totals["net_profit"]]]),
                ("Per-Bar Performance", ["Bar", "Revenue", "Cost", "Staff", "Net Profit", "Waste %"], [[r["bar_name"], r["revenue"], r["cost"], r["staff_cost"], r["net_profit"], r["waste_pct"]] for r in bar_rows]),
                ("Top Products", ["Product", "Units Sold", "Revenue"], [[r["product_name"], r["units_sold"], r["revenue"]] for r in product_rows]),
                ("Payroll", ["Employee", "Bar", "Salary"], [[r["full_name"], r["bar_name"], r["salary_paid"]] for r in payroll_rows]),
            ],
        ),
    )


def report_bundle(db: Session, event_id: int) -> ReportFile:
    rows = bars_summary(db, event_id)
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        for file_type in ("pdf", "xlsx"):
            all_bars = all_bars_report(db, event_id, file_type)
            full = full_event_report(db, event_id, file_type)
            archive.writestr(all_bars.filename, all_bars.data)
            archive.writestr(full.filename, full.data)
        for row in rows:
            bar_pdf = per_bar_report(db, row["bar_id"], "pdf")
            archive.writestr(bar_pdf.filename, bar_pdf.data)
    return ReportFile(f"event-{event_id}-report-bundle.zip", "application/zip", buffer.getvalue())
